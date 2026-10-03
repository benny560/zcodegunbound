# 提取 Boomer开炮表v4.xlsx 全部数据为结构化 JSON
import openpyxl, json, re

wb = openpyxl.load_workbook('Boomer开炮表v4.xlsx', data_only=True)
out = {}

# 1) 表1 无风力 度
ws = wb['表1_无风力 度']
nowind = []
for row in ws.iter_rows(min_row=4, values_only=True):
    if row[0] is None: continue
    nowind.append({'sd': int(row[0]), 'p81': row[1], 'p70': row[2], 'p50': row[3]})
out['nowind'] = nowind

# 2) 6个矩阵
def parse_cell(cell):
    """'1.137 / 82.1' -> {coef:1.137, angle:82.1, mode:'angle'}
       '0.250 / -0.011' -> {coef:0.250, dpower:-0.011, mode:'power'}
       '--' -> None
       None -> None"""
    if cell is None: return None
    s = str(cell).strip()
    if s in ('--','', 'None'): return None
    # 匹配 数字 / [-]数字
    m = re.match(r'^(-?[\d.]+)\s*/\s*(-?[\d.]+)$', s)
    if not m:
        # 可能是纯数字(角度)
        try:
            v = float(s)
            return {'raw': v}
        except: return None
    coef = float(m.group(1))
    second = float(m.group(2))
    if second < 0:
        return {'coef': coef, 'dpower': second, 'mode': 'power'}
    else:
        return {'coef': coef, 'angle': second, 'mode': 'angle'}

matrices = {}
matrix_sheets = [
    ('81_上', '矩阵_81°_上', 0),   # 风向 0-90
    ('81_下', '矩阵_81°_下', 270), # 风向 270-360
    ('70_上', '矩阵_70°_上', 0),
    ('70_下', '矩阵_70°_下', 270),
    ('50_上', '矩阵_50°_上', 0),
    ('50_下', '矩阵_50°_下', 270),
]

for key, sheet, base_wa in matrix_sheets:
    ws = wb[sheet]
    # 找表头行(含'风力\\风向')
    header_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row[0] and '风力' in str(row[0]):
            header_row = i
            break
    if header_row is None:
        print(f'WARN: {sheet} no header')
        continue
    rows = list(ws.iter_rows(values_only=True))
    wa_headers = rows[header_row-1]  # 风向列
    wa_list = [int(x) for x in wa_headers[1:] if x is not None]
    entries = []
    for row in rows[header_row:]:
        if row[0] is None: continue
        try:
            ws_val = int(row[0])
        except:
            continue
        for j, wa in enumerate(wa_list):
            cell = row[1+j]
            parsed = parse_cell(cell)
            if parsed is None: continue
            entries.append({'ws': ws_val, 'wa': wa, **parsed})
    matrices[key] = {
        'sheet': sheet, 'base_wa': base_wa,
        'wa_list': wa_list, 'entries': entries
    }

out['matrices'] = matrices

with open('_boomer_table_v4.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=1)

# 统计
print('提取完成 -> _boomer_table_v4.json')
print(f'无风表: {len(nowind)} 行 (sd 1-{nowind[-1]["sd"]})')
for k, v in matrices.items():
    angle_cnt = sum(1 for e in v['entries'] if e.get('mode')=='angle')
    power_cnt = sum(1 for e in v['entries'] if e.get('mode')=='power')
    print(f'  {k}: {len(v["entries"])} 条 (角度模式{angle_cnt}, 减力模式{power_cnt}), ws范围{min(e["ws"] for e in v["entries"])}-{max(e["ws"] for e in v["entries"])}, wa范围{v["wa_list"][0]}-{v["wa_list"][-1]}')
