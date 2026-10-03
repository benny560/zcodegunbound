# -*- coding: utf-8 -*-
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def probe(path):
    print("="*80)
    print("FILE:", path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        print("-"*70)
        print("SHEET:", ws.title, " dims:", ws.max_row, "x", ws.max_column)
        # header row(s)
        rows = list(ws.iter_rows(min_row=1, max_row=8, values_only=True))
        for i, r in enumerate(rows):
            print(f"  R{i+1}:", r)

for p in [
    r"C:\Users\fanli\Documents\zcodegunbound\实测数据v1.xlsx",
    r"C:\Users\fanli\Documents\zcodegunbound\识别数据.xlsx",
    r"C:\Users\fanli\Documents\zcodegunbound\gunbound calculator.xlsx",
]:
    try:
        probe(p)
    except Exception as e:
        print("ERR", p, e)
